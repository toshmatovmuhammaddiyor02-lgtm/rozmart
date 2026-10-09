"""Онлайн-платежи: отчёт для контроля кассы (раздел панели «Онлайн-платежи»).

Деньги онлайн-оплаты приходят на основной счёт сразу, через транзитный счёт
платёжной системы, уже без её комиссии. Отчёт показывает за период:

* каждый платёж — заказ, способ, сумма, комиссия, сколько придёт на счёт,
  возврат, фискальный чек, что с заказом на кассе REGOS;
* итоги по дням и способам — чтобы сверить с выпиской основного счёта;
* предупреждения — что требует внимания человека.

Платёж берётся из таблиц платёжных систем (Payme, Click, Paynet), а не из
заказа: у заказа один платёж, а у платёжной системы бывают и возвраты, и
деньги за заказ, который их уже не ждал. Оплата внутри Telegram — по заказу.

«Закрыт на кассе» — по статусу документа в REGOS («Завершен»). Наш статус
«Выполнен» ставит и доставщик, поэтому по нему закрытие на кассе не видно.
Статусы спрашиваем у REGOS при построении отчёта; нет связи — так и пишем.
"""
import io
import logging
import os
import time
from datetime import date, datetime, timedelta, timezone

from sqlalchemy import select
from sqlalchemy.orm import Session

from app import payments
from app.models import ClickTransaction, Order, PaymeTransaction, PaynetTransaction, utcnow

log = logging.getLogger(__name__)

TASHKENT = timezone(timedelta(hours=5))
TITLES = {"payme": "Payme", "click": "Click", "paynet": "Paynet", "telegram": "Telegram"}

# Через сколько часов оплаченный, но не закрытый на кассе заказ — повод
# посмотреть. Доставка «как можно скорее» — 60–90 минут, с запасом
OPEN_HOURS = float(os.getenv("REPORT_OPEN_HOURS", "3"))
# Сколько ждём ссылку на фискальный чек от Payme/Click, прежде чем тревожиться
FISCAL_WAIT = timedelta(minutes=30)
# Оплаченный заказ уходит в REGOS за секунды; дольше — что-то не так
PUSH_WAIT = timedelta(minutes=15)

# статусы документа REGOS: 26 «Завершен» — закрыт на кассе, 27 «Отменен»
REGOS_CLOSED, REGOS_CANCELED = 26, 27

ORDER_STATUS = {"NEW": "Ждёт оплаты", "CONFIRMED": "Принят", "ASSEMBLING": "Собирается",
                "DELIVERING": "В доставке", "DONE": "Выполнен", "CANCELED": "Отменён"}


# ------------------------------------------------------------ период

def parse_day(value: str | None, default: date) -> date:
    try:
        return date.fromisoformat(value) if value else default
    except ValueError:
        return default


def today() -> date:
    return datetime.now(TASHKENT).date()


def bounds(day_from: date, day_to: date) -> tuple[datetime, datetime]:
    """Сутки по Ташкенту -> UTC без пояса, как в базе. Конец не включается."""
    start = datetime.combine(day_from, datetime.min.time(), TASHKENT)
    end = datetime.combine(day_to + timedelta(days=1), datetime.min.time(), TASHKENT)
    return (start.astimezone(timezone.utc).replace(tzinfo=None),
            end.astimezone(timezone.utc).replace(tzinfo=None))


def local(moment: datetime | None) -> datetime | None:
    if moment is None:
        return None
    return moment.replace(tzinfo=timezone.utc).astimezone(TASHKENT)


def from_ms(ms: int | None) -> datetime | None:
    return datetime.fromtimestamp(ms / 1000, timezone.utc).replace(tzinfo=None) if ms else None


# ------------------------------------------------------------ комиссия

def fee_percent(provider: str) -> float:
    payments.refresh()
    raw = str(getattr(payments, f"FEE_{provider.upper()}", "") or "0").replace(",", ".")
    try:
        return float(raw)
    except ValueError:
        return 0.0


# ------------------------------------------------------------ платежи

def _row(provider, paid_at, order: Order | None, number, amount, state, state_text,
         fiscal, fiscal_note="", refunded_at=None):
    pct = fee_percent(provider)
    refunded = state in ("refunded", "refund_failed")
    fee = 0 if refunded else round(amount * pct / 100)
    return {
        "provider": provider, "title": TITLES.get(provider, provider),
        "paid_at": paid_at, "order": order, "number": number,
        "customer": order.customer_name if order else "",
        "amount": amount, "fee_pct": pct, "fee": fee,
        "net": 0 if refunded else amount - fee,
        "state": state, "state_text": state_text, "refunded_at": refunded_at,
        "fiscal": fiscal, "fiscal_note": fiscal_note,
    }


def _fiscal(url: str | None, paid_at: datetime, note: str = "") -> tuple[str, str]:
    if url:
        return "ok", url
    if note:
        return "error", note
    return ("wait" if utcnow() - paid_at < FISCAL_WAIT else "missing"), ""


def collect(db: Session, day_from: date, day_to: date) -> list[dict]:
    start, end = bounds(day_from, day_to)
    rows: list[dict] = []
    orders: dict[int, Order] = {}

    def order_of(order_id):
        if order_id is None:
            return None
        if order_id not in orders:
            orders[order_id] = db.get(Order, order_id)
        return orders[order_id]

    # Payme: проведённые (2) и возвращённые после проведения (-2)
    start_ms = int(start.replace(tzinfo=timezone.utc).timestamp() * 1000)
    end_ms = int(end.replace(tzinfo=timezone.utc).timestamp() * 1000)
    for tx in db.scalars(select(PaymeTransaction).where(
            PaymeTransaction.state.in_((2, -2)),
            PaymeTransaction.perform_time >= start_ms, PaymeTransaction.perform_time < end_ms)):
        order = order_of(tx.order_id)
        paid_at = from_ms(tx.perform_time)
        url = tx.fiscal_qr_url or (order.fiscal_qr_url if order else None)
        fiscal, note = _fiscal(url, paid_at)
        if tx.state == -2:
            rows.append(_row("payme", paid_at, order, tx.order_number, tx.amount // payments.MINOR,
                             "refunded", "Возврат через Payme", fiscal, note,
                             refunded_at=from_ms(tx.cancel_time)))
        else:
            rows.append(_row("payme", paid_at, order, tx.order_number, tx.amount // payments.MINOR,
                             "paid", "Оплачен", fiscal, note))

    # Click: проведённые и деньги за заказ, который их уже не ждал
    from app import click
    for tx in db.scalars(select(ClickTransaction).where(
            ClickTransaction.state.in_((click.COMPLETED, click.REVERSED, click.REVERSE_FAILED)),
            ClickTransaction.completed_at >= start, ClickTransaction.completed_at < end)):
        order = order_of(tx.order_id)
        url = order.fiscal_qr_url if order else None
        error = ""
        if tx.fiscal_status == "error" and not url:
            error = "чек не принят: " + (tx.fiscal_note or "")[:200]
        fiscal, note = _fiscal(url, tx.completed_at, error)
        if tx.state == click.REVERSED:
            rows.append(_row("click", tx.completed_at, order, tx.order_number, tx.amount,
                             "refunded", "Возвращён: заказ уже не ждал оплату", "none", "",
                             refunded_at=tx.canceled_at))
        elif tx.state == click.REVERSE_FAILED:
            rows.append(_row("click", tx.completed_at, order, tx.order_number, tx.amount,
                             "refund_failed", "Вернуть вручную в кабинете Click", "none", "",
                             refunded_at=None))
        else:
            rows.append(_row("click", tx.completed_at, order, tx.order_number, tx.amount,
                             "paid", "Оплачен", fiscal, note))

    # Paynet: проведённые и отменённые Paynet после проведения
    for tx in db.scalars(select(PaynetTransaction).where(
            PaynetTransaction.created_at >= start, PaynetTransaction.created_at < end)):
        if tx.order_number == "TEST-PAYNET":
            continue                        # тестовый счёт Paynet — не деньги
        order = order_of(tx.order_id)
        refunded = tx.state == 2
        rows.append(_row("paynet", tx.created_at, order, tx.order_number,
                         tx.amount // payments.MINOR,
                         "refunded" if refunded else "paid",
                         "Paynet отменил оплату" if refunded else "Оплачен", "none", "",
                         refunded_at=tx.canceled_at if refunded else None))

    # Оплата внутри Telegram — у неё нет своей таблицы, только заказ
    for order in db.scalars(select(Order).where(
            Order.paid_at >= start, Order.paid_at < end,
            Order.payment_charge_id.is_not(None),
            ~Order.payment_charge_id.like("payme:%"), ~Order.payment_charge_id.like("click:%"),
            ~Order.payment_charge_id.like("paynet:%"))):
        rows.append(_row("telegram", order.paid_at, order, order.number,
                         order.paid_amount or order.total, "paid", "Оплачен", "none", ""))

    rows.sort(key=lambda r: r["paid_at"], reverse=True)
    attach_regos(rows)
    return rows


# ------------------------------------------------------------ касса REGOS

_regos_cache: dict = {"at": 0.0, "key": None, "found": {}}


def attach_regos(rows: list[dict]) -> None:
    """Статус документа в REGOS к каждому платежу. Нет связи — regos_error."""
    ids = sorted({r["order"].regos_document_id for r in rows
                  if r["order"] is not None and r["order"].regos_document_id})
    found, error = {}, ""
    if ids:
        key = tuple(ids)
        if _regos_cache["key"] == key and time.time() - _regos_cache["at"] < 60:
            found = _regos_cache["found"]
        else:
            try:
                from app.regos.client import RegosClient
                from app.regos.orders_sync import fetch_statuses
                found = fetch_statuses(RegosClient(), ids)
                _regos_cache.update(at=time.time(), key=key, found=found)
            except Exception as exc:            # noqa: BLE001
                log.warning("Отчёт: статусы REGOS не получены: %s", exc)
                error = "нет связи с REGOS"
    for r in rows:
        order = r["order"]
        doc = order.regos_document_id if order else None
        status = found.get(doc) if doc else None
        r["regos_doc"] = doc
        r["regos_status_id"] = status.get("id") if status else None
        r["regos_status"] = (status.get("name") if status else
                             (error if doc else ("не выгружен" if order else "")))


# ------------------------------------------------------------ предупреждения

def alerts(rows: list[dict]) -> list[dict]:
    now = utcnow()
    found = []

    def add(row, kind, text):
        found.append({"kind": kind, "number": row["number"], "provider": row["title"],
                      "text": text})

    for r in rows:
        order = r["order"]
        if r["state"] == "refund_failed":
            add(r, "refund", "Деньги пришли за заказ, который их уже не ждал, и автоматически "
                             "не вернулись — верните в кабинете Click")
        if r["state"] != "paid" or order is None:
            continue
        if order.status == "CANCELED":
            add(r, "refund", "Заказ отменён, а оплата не возвращена — верните деньги в кабинете "
                             + r["title"] + " или восстановите заказ")
            continue
        if not r["regos_doc"] and now - r["paid_at"] > PUSH_WAIT:
            add(r, "regos", "Оплачен, но не выгружен в REGOS"
                + (f": {order.regos_error}" if order.regos_error else ""))
        elif (r["regos_status_id"] and r["regos_status_id"] not in (REGOS_CLOSED, REGOS_CANCELED)
              and now - r["paid_at"] > timedelta(hours=OPEN_HOURS)):
            hours = int((now - r["paid_at"]).total_seconds() // 3600)
            add(r, "cash", f"Оплачен {hours} ч назад, а на кассе не закрыт "
                           f"(в REGOS: {r['regos_status']})")
        if r["fiscal"] == "missing":
            add(r, "fiscal", "Нет ссылки на фискальный чек от " + r["title"])
        elif r["fiscal"] == "error":
            add(r, "fiscal", "Фискальный чек: " + r["fiscal_note"])
        if order.fiscal_qr_url and order.fiscal_qr_regos_at is None and order.fiscal_qr_regos_error:
            add(r, "fiscal", "Ссылка на чек не передана в REGOS: " + order.fiscal_qr_regos_error[:200])
    return found


# ------------------------------------------------------------ итоги

def totals(rows: list[dict]) -> dict:
    """Итоги по дням и способам и за весь период."""
    days: dict[tuple, dict] = {}
    by_provider: dict[str, dict] = {}
    overall = _empty("Всего")
    for r in rows:
        day = local(r["paid_at"]).date().isoformat()
        for bucket in (days.setdefault((day, r["provider"]), _empty(r["title"], day)),
                       by_provider.setdefault(r["provider"], _empty(r["title"])), overall):
            bucket["count"] += 1
            bucket["paid"] += r["amount"]
            if r["state"] in ("refunded", "refund_failed"):
                bucket["refunds"] += r["amount"]
                bucket["refund_count"] += 1
            bucket["fee"] += r["fee"]
            bucket["net"] += r["net"]
    return {"days": sorted(days.values(), key=lambda d: (d["day"], d["title"]), reverse=True),
            "providers": list(by_provider.values()), "overall": overall}


def _empty(title, day=None):
    return {"title": title, "day": day, "count": 0, "paid": 0, "refunds": 0,
            "refund_count": 0, "fee": 0, "net": 0}


# ------------------------------------------------------------ ответ панели

def view(db: Session, day_from: date, day_to: date) -> dict:
    rows = collect(db, day_from, day_to)
    return {
        "from": day_from.isoformat(), "to": day_to.isoformat(),
        "fees": {p: fee_percent(p) for p in ("payme", "click", "paynet")},
        "open_hours": OPEN_HOURS,
        "alerts": alerts(rows),
        "totals": totals(rows),
        "payments": [{
            "time": local(r["paid_at"]).strftime("%d.%m.%Y %H:%M"),
            "number": r["number"], "customer": r["customer"], "provider": r["title"],
            "amount": r["amount"], "fee_pct": r["fee_pct"], "fee": r["fee"], "net": r["net"],
            "state": r["state"], "state_text": r["state_text"],
            "fiscal": r["fiscal"], "fiscal_url": r["fiscal_note"] if r["fiscal"] == "ok" else "",
            "fiscal_note": r["fiscal_note"] if r["fiscal"] != "ok" else "",
            "regos_doc": r["regos_doc"], "regos_status": r["regos_status"],
            "order_status": ORDER_STATUS.get(r["order"].status, r["order"].status)
            if r["order"] else "",
        } for r in rows],
    }


FISCAL_TEXT = {"ok": "есть", "wait": "ждём", "missing": "нет", "error": "ошибка", "none": "—"}


def xlsx(db: Session, day_from: date, day_to: date) -> bytes:
    """Excel: «Платежи» и «Итоги по дням»."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    rows = collect(db, day_from, day_to)
    book = Workbook()
    sheet = book.active
    sheet.title = "Платежи"
    head = ["Дата", "Время", "Заказ", "Покупатель", "Способ", "Сумма", "Комиссия, %",
            "Комиссия", "На счёт", "Состояние", "Возврат", "Фискальный чек",
            "Документ REGOS", "Статус в REGOS", "Статус заказа"]
    sheet.append(head)
    for r in rows:
        moment = local(r["paid_at"])
        refund = r["amount"] if r["state"] in ("refunded", "refund_failed") else 0
        fiscal = r["fiscal_note"] if r["fiscal"] == "ok" else FISCAL_TEXT.get(r["fiscal"], "")
        sheet.append([moment.strftime("%d.%m.%Y"), moment.strftime("%H:%M"), r["number"],
                      r["customer"], r["title"], r["amount"], r["fee_pct"], r["fee"], r["net"],
                      r["state_text"], refund, fiscal, r["regos_doc"] or "", r["regos_status"],
                      ORDER_STATUS.get(r["order"].status, "") if r["order"] else ""])
    _style(sheet, head, Font, get_column_letter, money_cols=(6, 8, 9, 11))

    days = book.create_sheet("Итоги по дням")
    head2 = ["Дата", "Способ", "Платежей", "Оплачено", "Возвратов", "Возвращено",
             "Комиссия", "Должно прийти на счёт"]
    days.append(head2)
    summary = totals(rows)
    for d in summary["days"]:
        days.append([date.fromisoformat(d["day"]).strftime("%d.%m.%Y"), d["title"], d["count"],
                     d["paid"], d["refund_count"], d["refunds"], d["fee"], d["net"]])
    o = summary["overall"]
    days.append(["Итого", "", o["count"], o["paid"], o["refund_count"], o["refunds"], o["fee"],
                 o["net"]])
    for cell in days[days.max_row]:
        cell.font = Font(bold=True)
    _style(days, head2, Font, get_column_letter, money_cols=(4, 6, 7, 8))

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _style(sheet, head, Font, get_column_letter, money_cols=()):
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for i, title in enumerate(head, start=1):
        width = max([len(str(title))] + [len(str(c.value or "")) for c in sheet[get_column_letter(i)]])
        sheet.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 9), 60)
        if i in money_cols:
            for c in sheet[get_column_letter(i)][1:]:
                c.number_format = "# ##0"
