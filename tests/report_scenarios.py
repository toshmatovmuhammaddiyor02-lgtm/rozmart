"""Отчёт «Онлайн-платежи» в панели — через HTTP к приложению.

    .venv\\Scripts\\python.exe tests/report_scenarios.py

Платежи Payme и Click за период, комиссия и сумма на счёт, возвраты,
предупреждения (не закрыт на кассе, нет чека, заказ отменён — деньги не
возвращены), итоги по дням, выгрузка в Excel, доступ только администратору.
База — временный SQLite, REGOS и Click подменены. Выход с кодом 1, если
хоть одна проверка не прошла.
"""
import base64
import hashlib
import hmac
import io
import json
import os
import pathlib
import sys
import tempfile
import time
from datetime import timedelta
from urllib.parse import urlencode

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

tmp = pathlib.Path(tempfile.mkdtemp()) / "t.db"
os.environ.update({
    "DATABASE_URL": f"sqlite:///{tmp.as_posix()}",
    "BOT_TOKEN": "111:shop-token",
    "ADMIN_TELEGRAM_IDS": "1",
    "REGOS_KEY": "test-key",
    "REGOS_STOCK_ID": "5",
    "REGOS_PRICE_TYPE_ID": "5",
    "REGOS_ORDER_FROM_ID": "2",
    "NOTIFY_BOT_TOKEN": "222:staff-token",
    "PAYMENT_PROVIDERS": "payme,click",
    "PAYME_MERCHANT_ID": "65f0c0ffee0000000000abcd",
    "PAYME_KEY": "sandbox-key-123",
    "PAYME_TEST": "0",
    "CLICK_SERVICE_ID": "113234",
    "CLICK_MERCHANT_ID": "50001",
    "CLICK_MERCHANT_USER_ID": "70001",
    "CLICK_SECRET_KEY": "Cl1ck-Secret-Key",
    "CLICK_TIN": "312345678",
    "PAYMENT_TEST_USERS": "900",
})

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import app.regos.orders_push as orders_push  # noqa: E402
from app import click, notify  # noqa: E402
from app.db import SessionLocal, init_db  # noqa: E402
from app.main import app  # noqa: E402
from app.models import Category, Order, PaymeTransaction, Product, Setting, Variant, utcnow  # noqa: E402

notify.call = lambda method, payload, token=None: {"message_id": 1}

# REGOS: документы по порядку; статус документа задаём сами
doc_status = {}


class FakeRegos:
    def call(self, method, params=None, **kw):
        if method == "DocOrderDelivery/Get":
            if params and params.get("filters") and params["filters"][0].get("Field") == "id":
                ids = [int(x) for x in str(params["filters"][0]["Value"]).split(",")]
                return [{"id": i, "status": doc_status.get(i, {"id": 24, "name": "В обработке"})}
                        for i in ids]
            return []
        if method == "RetailCustomer/Get":
            return [{"id": 77}]
        if method == "DocOrderDelivery/AddFull":
            FakeRegos.n = getattr(FakeRegos, "n", 100) + 1
            return {"new_id": FakeRegos.n}
        return {"row_affected": 1}


orders_push.RegosClient = FakeRegos
import app.regos.client as regos_client  # noqa: E402
regos_client.RegosClient = FakeRegos
click.api = lambda method, path, body=None: (
    {"error_code": 0} if "submit_items" in path else {"error_code": -1})

init_db()
db = SessionLocal()
cat = Category(name="Колбасы"); db.add(cat); db.flush()
prod = Product(category_id=cat.id, name="Докторская", is_active=True); db.add(prod); db.flush()
variant = Variant(product_id=prod.id, external_code="003100", weight="400 г", price=40000,
                  regos_price=40000, is_active=True, mxik="01601002002011001",
                  package_code="1403931", vat_percent=12)
db.add(variant)
db.add(Setting(name=orders_push.SETTING_AUTO, value="1"))
db.add(Setting(name="delivery_price", value="0"))
db.commit()
client = TestClient(app)


def init_data(user_id):
    user = json.dumps({"id": user_id, "first_name": "Test"})
    fields = {"auth_date": str(int(time.time())), "user": user}
    check = "\n".join(f"{k}={fields[k]}" for k in sorted(fields))
    secret = hmac.new(b"WebAppData", b"111:shop-token", hashlib.sha256).digest()
    fields["hash"] = hmac.new(secret, check.encode(), hashlib.sha256).hexdigest()
    return {"X-Telegram-Init-Data": urlencode(fields)}


BUYER, ADMIN = init_data(900), init_data(1)
PAYME = {"Authorization": "Basic " + base64.b64encode(b"Paycom:sandbox-key-123").decode()}
results = []
seq = [0]


def check(n, name, ok, detail=""):
    results.append((n, name, ok))
    print(f"{'ПРОШЛО ' if ok else 'НЕ ПРОШЛО'} {n:>2}. {name}{(' — ' + str(detail)) if detail else ''}")


def new_order(key, provider, qty=1):
    r = client.post("/api/orders", json={
        "customer_name": "Javohir Rahmonov", "phone": "+998901234567", "address": "Chilonzor 9",
        "delivery_slot": "Как можно скорее", "payment_method": "online", "consent": True,
        "client_key": key, "pay_provider": provider,
        "items": [{"variant_id": variant.id, "quantity": qty}]}, headers=BUYER)
    assert r.status_code == 201, r.text
    return r.json()


def rpc(method, params):
    seq[0] += 1
    return client.post("/payme", json={"jsonrpc": "2.0", "id": seq[0], "method": method,
                                       "params": params}, headers=PAYME).json()


def pay_payme(order, tx_id):
    rpc("CreateTransaction", {"id": tx_id, "time": int(time.time() * 1000),
                              "amount": order["total"] * 100, "account": {"order_id": order["number"]}})
    return rpc("PerformTransaction", {"id": tx_id})


def click_form(click_id, number, amount, action, prepare_id=None):
    data = {"click_trans_id": str(click_id), "service_id": "113234",
            "click_paydoc_id": str(click_id + 5000000), "merchant_trans_id": number,
            "amount": amount, "action": str(action), "error": "0", "error_note": "Success",
            "sign_time": "2026-10-07 12:00:00"}
    parts = [data["click_trans_id"], "113234", "Cl1ck-Secret-Key", number]
    if prepare_id is not None:
        data["merchant_prepare_id"] = str(prepare_id)
        parts.append(str(prepare_id))
    parts += [amount, str(action), data["sign_time"]]
    data["sign_string"] = hashlib.md5("".join(parts).encode()).hexdigest()
    return data


def pay_click(order, click_id):
    amount = f"{order['total']}.00"
    r = client.post("/click/prepare", data=click_form(click_id, order["number"], amount, 0)).json()
    return client.post("/click/complete", data=click_form(click_id, order["number"], amount, 1,
                                                           r["merchant_prepare_id"])).json()


def shift(model, where, **values):
    s = SessionLocal()
    s.query(model).filter_by(**where).update(values)
    s.commit()
    s.close()


def report(**params):
    r = client.get("/api/admin/online-payments", params=params, headers=ADMIN)
    assert r.status_code == 200, r.text
    return r.json()


# ------------------------------------------------------------ данные
p1 = new_order("rep-key-001", "payme")            # 40 000 — оплачен, закрыт на кассе
pay_payme(p1, "tx-1")
p2 = new_order("rep-key-002", "payme", qty=2)     # 80 000 — оплачен, затем возврат
pay_payme(p2, "tx-2")
rpc("CancelTransaction", {"id": "tx-2", "reason": 5})
c1 = new_order("rep-key-003", "click", qty=3)     # 120 000 — Click
pay_click(c1, 9001)
y1 = new_order("rep-key-004", "payme")            # вчерашний
pay_payme(y1, "tx-4")
yesterday = int((time.time() - 86400) * 1000)
shift(PaymeTransaction, {"payme_id": "tx-4"}, perform_time=yesterday)

s = SessionLocal()
d1 = s.query(Order).filter_by(number=p1["number"]).one().regos_document_id
s.close()
doc_status[d1] = {"id": 26, "name": "Завершен"}

# ------------------------------------------------------------ отчёт за сегодня
r = report()
nums = [p["number"] for p in r["payments"]]
check(1, "Сегодня: три платежа, вчерашний не попал",
      sorted(nums) == sorted([p1["number"], p2["number"], c1["number"]]), nums)
pm = {t["title"]: t for t in r["totals"]["providers"]}
check(2, "Payme: 2 платежа на 120 000, возврат 80 000, комиссия 1% с оставшегося",
      pm["Payme"]["count"] == 2 and pm["Payme"]["paid"] == 120000
      and pm["Payme"]["refunds"] == 80000 and pm["Payme"]["fee"] == 400
      and pm["Payme"]["net"] == 39600, pm["Payme"])
check(2, "Click: 120 000, комиссия 2% — на счёт 117 600",
      pm["Click"]["fee"] == 2400 and pm["Click"]["net"] == 117600, pm["Click"])
check(2, "Всего на счёт 157 200", r["totals"]["overall"]["net"] == 157200, r["totals"]["overall"])
row2 = next(p for p in r["payments"] if p["number"] == p2["number"])
check(3, "Возврат виден в строке платежа", row2["state"] == "refunded"
      and "Возврат" in row2["state_text"] and row2["net"] == 0)
row1 = next(p for p in r["payments"] if p["number"] == p1["number"])
check(3, "Касса REGOS: «Завершен» с номером документа", row1["regos_status"] == "Завершен"
      and row1["regos_doc"] == d1, row1)
check(3, "Чек Payme ещё ждём — первые 30 минут без тревоги", row1["fiscal"] == "wait"
      and not any(a["kind"] == "fiscal" for a in r["alerts"]))

# ------------------------------------------------------------ предупреждения
shift(PaymeTransaction, {"payme_id": "tx-1"}, perform_time=int((time.time() - 3600) * 1000))
shift(Order, {"number": c1["number"]}, paid_at=utcnow() - timedelta(hours=4))
s = SessionLocal()
s.query(click.ClickTransaction).filter_by(click_trans_id=9001).update(
    {"completed_at": utcnow() - timedelta(hours=4)})
s.commit(); s.close()
r = report()
kinds = {(a["number"], a["kind"]) for a in r["alerts"]}
check(4, "Click оплачен 4 ч назад, на кассе «В обработке» — предупреждение",
      (c1["number"], "cash") in kinds, r["alerts"])
check(4, "Закрытый на кассе заказ не тревожит", (p1["number"], "cash") not in kinds)
check(4, "Чека Payme нет больше 30 минут — предупреждение", (p1["number"], "fiscal") in kinds)
check(4, "Возвращённый платёж предупреждений не даёт", not any(a["number"] == p2["number"]
                                                              for a in r["alerts"]))
p5 = new_order("rep-key-005", "payme")
pay_payme(p5, "tx-5")
shift(Order, {"number": p5["number"]}, status="CANCELED")
r = report()
check(5, "Заказ отменён, а деньги не возвращены — предупреждение",
      any(a["number"] == p5["number"] and a["kind"] == "refund" for a in r["alerts"]))

# ------------------------------------------------------------ период, комиссия, доступ
r = report(date_from=(utcnow() - timedelta(days=1)).strftime("%Y-%m-%d"))
check(6, "Два дня — и вчерашний платёж, итоги по дням раздельно",
      y1["number"] in [p["number"] for p in r["payments"]]
      and len({d["day"] for d in r["totals"]["days"]}) == 2, [d["day"] for d in r["totals"]["days"]])
cfg = client.put("/api/admin/payments/config", json={"values": {"FEE_CLICK": "abc"}}, headers=ADMIN)
check(7, "Комиссия — только число", cfg.status_code == 400)
cfg = client.put("/api/admin/payments/config", json={"values": {"FEE_CLICK": "1,5"}}, headers=ADMIN)
r = report()
ck = {t["title"]: t for t in r["totals"]["providers"]}["Click"]
check(7, "Комиссию меняют в панели: Click 1,5% — на счёт 118 200",
      cfg.status_code == 200 and ck["fee"] == 1800 and ck["net"] == 118200, ck)
check(8, "Покупателю отчёт недоступен",
      client.get("/api/admin/online-payments", headers=BUYER).status_code in (401, 403, 404))

# ------------------------------------------------------------ Excel
x = client.get("/api/admin/online-payments.xlsx", headers=ADMIN)
book = load_workbook(io.BytesIO(x.content)) if x.status_code == 200 else None
check(9, "Excel скачивается", x.status_code == 200 and "attachment" in x.headers.get(
    "content-disposition", ""), x.status_code)
if book:
    sheet = book["Платежи"]
    head = [c.value for c in sheet[1]]
    check(9, "Лист «Платежи»: столбцы как договорились", head[:6] == [
        "Дата", "Время", "Заказ", "Покупатель", "Способ", "Сумма"] and "Фискальный чек" in head
        and "Документ REGOS" in head, head)
    check(9, "Строк столько же, сколько платежей", sheet.max_row - 1 == len(report()["payments"]))
    days = book["Итоги по дням"]
    check(9, "Лист «Итоги по дням» с итоговой строкой", days.cell(days.max_row, 1).value == "Итого")

failed = [r for r in results if not r[2]]
print(f"\nИтого проверок: {len(results)}, не прошло: {len(failed)}")
sys.exit(1 if failed else 0)
